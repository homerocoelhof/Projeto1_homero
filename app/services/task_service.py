"""Business rules for summarizing PDFs and generating Word documents."""

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from threading import Lock
from typing import Protocol

import pymupdf
from docx import Document

from app.models.task import SummaryMode, TaskCreate, TaskOut, TaskUpdate
from app.repositories.task_repository import TaskRepository


class PriorityAdvisor(Protocol):
    def summarize(self, source_text: str, sentence: str | None) -> str: ...


@dataclass(frozen=True)
class TaskResult:
    task: TaskOut
    document: bytes
    filename: str


class TaskProcessingError(Exception):
    """Controlled error during task processing."""


class TaskService:
    """Coordinates PDF extraction, summary generation, and DOCX creation."""

    def __init__(self, repository: TaskRepository, priority_advisor: PriorityAdvisor) -> None:
        self._repository = repository
        self._priority_advisor = priority_advisor
        self._filename_lock = Lock()
        self._filename_counters: dict[str, int] = {}

    def create_summary(
        self,
        pdf_path: str | Path,
        sentence: str | None = None,
        *,
        summary_mode: SummaryMode = "guided",
    ) -> TaskResult:
        """Create a sentence-guided or full-document summary as a DOCX."""
        path = Path(pdf_path).expanduser()
        normalized_sentence = self._validate_summary_request(sentence, summary_mode)
        if path.suffix.lower() != ".pdf":
            raise ValueError("O arquivo informado deve ter extensão .pdf.")
        if not path.is_file():
            raise ValueError("O arquivo PDF informado não existe ou não é um arquivo.")

        return self._process_summary(
            pdf_source=path,
            source_filename=path.name,
            sentence=normalized_sentence,
            summary_mode=summary_mode,
        )

    def create_summary_from_bytes(
        self,
        pdf_content: bytes,
        source_filename: str,
        sentence: str | None = None,
        *,
        summary_mode: SummaryMode = "guided",
    ) -> TaskResult:
        """Create a summary directly from uploaded PDF bytes."""
        safe_filename = Path(source_filename).name
        if Path(safe_filename).suffix.lower() != ".pdf":
            raise ValueError("O arquivo informado deve ter extensão .pdf.")
        if not pdf_content:
            raise ValueError("O arquivo PDF está vazio.")

        normalized_sentence = self._validate_summary_request(sentence, summary_mode)
        return self._process_summary(
            pdf_source=pdf_content,
            source_filename=safe_filename,
            sentence=normalized_sentence,
            summary_mode=summary_mode,
        )

    def _process_summary(
        self,
        *,
        pdf_source: Path | bytes,
        source_filename: str,
        sentence: str | None,
        summary_mode: SummaryMode,
    ) -> TaskResult:
        task = self._repository.create(
            TaskCreate(
                source_filename=source_filename,
                summary_mode=summary_mode,
                sentence=sentence if summary_mode == "guided" else None,
            )
        )
        self._repository.update(task.id, TaskUpdate(status="processing"))

        try:
            paragraphs = self._extract_paragraphs(pdf_source)
            if summary_mode == "guided":
                normalized_query = self._normalize_for_search(normalized_sentence or "")
                relevant_paragraphs = [
                    paragraph
                    for paragraph in paragraphs
                    if normalized_query in self._normalize_for_search(paragraph)
                ]
                if not relevant_paragraphs:
                    raise TaskProcessingError(
                        f"A sentença não foi encontrada no arquivo '{source_filename}'."
                    )
                source_text = "\n\n".join(relevant_paragraphs)
            else:
                source_text = "\n\n".join(paragraphs)

            summary = self._priority_advisor.summarize(
                source_text=source_text,
                sentence=sentence if summary_mode == "guided" else None,
            ).strip()
            if not summary:
                raise TaskProcessingError("O componente de resumo retornou conteúdo vazio.")

            document = self._create_docx(
                normalized_sentence or "Documento inteiro",
                summary,
                source_filename=source_filename,
                heading="Sentença orientadora" if summary_mode == "guided" else "Escopo do resumo",
            )
            completed = self._repository.update(
                task.id,
                TaskUpdate(status="completed", summary=summary),
            )
            if completed is None:
                raise TaskProcessingError("A tarefa deixou de existir durante o processamento.")

            return TaskResult(
                task=completed,
                document=document,
                filename=self._next_output_filename(Path(source_filename).stem),
            )
        except Exception as exc:
            self._repository.update(task.id, TaskUpdate(status="failed"))
            if isinstance(exc, TaskProcessingError):
                raise
            raise TaskProcessingError("Não foi possível processar o arquivo PDF.") from exc

    @staticmethod
    def _validate_summary_request(
        sentence: str | None,
        summary_mode: SummaryMode,
    ) -> str | None:
        normalized_sentence = sentence.strip() if sentence else None
        if summary_mode not in {"guided", "whole"}:
            raise ValueError("Modo de resumo inválido.")
        if summary_mode == "guided" and not normalized_sentence:
            raise ValueError("A sentença orientadora não pode estar vazia.")
        return normalized_sentence

    @staticmethod
    def _normalize_for_search(text: str) -> str:
        return " ".join(text.casefold().split())

    def _next_output_filename(self, source_stem: str) -> str:
        key = source_stem.casefold()
        with self._filename_lock:
            next_number = self._filename_counters.get(key, 0) + 1
            self._filename_counters[key] = next_number
        return f"{source_stem}_{next_number}.docx"

    @staticmethod
    def _extract_text(pdf_source: Path | bytes) -> str:
        return "\n\n".join(TaskService._extract_paragraphs(pdf_source))

    @staticmethod
    def _extract_paragraphs(pdf_source: Path | bytes) -> list[str]:
        try:
            if isinstance(pdf_source, bytes):
                pdf_document = pymupdf.open(stream=pdf_source, filetype="pdf")
            else:
                pdf_document = pymupdf.open(pdf_source)
            with pdf_document as pdf:
                paragraphs = [
                    block[4].strip()
                    for page in pdf
                    for block in page.get_text("blocks")
                    if block[4].strip()
                ]
        except Exception as exc:
            raise TaskProcessingError("O arquivo não é um PDF válido ou não pôde ser lido.") from exc
        if not paragraphs:
            raise TaskProcessingError("O PDF não contém texto extraível.")
        return paragraphs

    @staticmethod
    def _create_docx(
        scope: str,
        summary: str,
        *,
        source_filename: str,
        heading: str,
    ) -> bytes:
        document = Document()
        document.add_heading("Resumo do PDF", level=1)
        document.add_paragraph(f"Arquivo analisado: {source_filename}")
        document.add_heading(heading, level=2)
        document.add_paragraph(scope)
        document.add_heading("Resumo", level=2)
        for line in summary.splitlines():
            content = line.strip()
            if not content:
                continue
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", content)
            numbered_item = re.match(r"^\d+[.)]\s+(.+)$", content)
            if heading_match:
                level = min(4, max(2, len(heading_match.group(1)) + 1))
                document.add_heading(heading_match.group(2).strip(), level=level)
            elif content.startswith("> "):
                paragraph = document.add_paragraph(style="Quote")
                TaskService._add_markdown_text(paragraph, content[2:].strip())
            elif content.startswith(("- ", "* ", chr(0x2022) + " ")):
                paragraph = document.add_paragraph(style="List Bullet")
                TaskService._add_markdown_text(paragraph, content[2:].strip())
            elif numbered_item:
                paragraph = document.add_paragraph(style="List Number")
                TaskService._add_markdown_text(paragraph, numbered_item.group(1).strip())
            else:
                paragraph = document.add_paragraph()
                TaskService._add_markdown_text(paragraph, content)
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _add_markdown_text(paragraph: object, content: str) -> None:
        """Preserve bold emphasis used for key findings in the generated Word file."""
        parts = re.split(r"(\*\*.*?\*\*)", content)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                paragraph.add_run(part[2:-2]).bold = True
            else:
                paragraph.add_run(part)
