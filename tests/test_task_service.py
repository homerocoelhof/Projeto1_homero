from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.task_routes import router
from app.repositories.task_repository import TaskRepository
from app.services.task_service import TaskProcessingError, TaskService


class FakePriorityAdvisor:
    """Substitui o resumidor externo por uma resposta previsível."""

    def summarize(self, source_text: str, sentence: str | None) -> str:
        assert "Texto de exemplo extraído do PDF." in source_text
        if sentence is None:
            return "Resumo completo de exemplo."
        assert sentence == "Qual é o assunto principal?"
        return "Este é um resumo de exemplo."


@pytest.fixture
def pdf_path(tmp_path: Path) -> Path:
    """Cria um PDF digital pequeno para os testes."""
    path = tmp_path / "documento.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text(
            (72, 72),
            "Texto de exemplo extraído do PDF. Qual é o assunto principal?",
        )
        pdf.save(path)
    return path


@pytest.fixture
def repository() -> TaskRepository:
    return TaskRepository()


@pytest.fixture
def service(repository: TaskRepository) -> TaskService:
    return TaskService(
        repository=repository,
        priority_advisor=FakePriorityAdvisor(),
    )


def test_create_summary_extracts_pdf_and_generates_word_document(
    service: TaskService,
    repository: TaskRepository,
    pdf_path: Path,
) -> None:
    result = service.create_summary(pdf_path, "Qual é o assunto principal?")

    assert result.task.status == "completed"
    assert result.task.source_filename == "documento.pdf"
    assert result.task.summary == "Este é um resumo de exemplo."
    assert result.filename == "documento_1.docx"
    assert repository.get_by_id(result.task.id) == result.task

    document = Document(BytesIO(result.document))
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Qual é o assunto principal?" in document_text
    assert "Este é um resumo de exemplo." in document_text


def test_create_whole_pdf_summary_from_bytes_without_guiding_sentence(
    service: TaskService,
    repository: TaskRepository,
    pdf_path: Path,
) -> None:
    result = service.create_summary_from_bytes(
        pdf_content=pdf_path.read_bytes(),
        source_filename=pdf_path.name,
        summary_mode="whole",
    )

    assert result.task.status == "completed"
    assert result.task.summary_mode == "whole"
    assert result.task.sentence is None
    assert result.task.summary == "Resumo completo de exemplo."
    assert repository.get_by_id(result.task.id) == result.task

    document = Document(BytesIO(result.document))
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Documento inteiro" in document_text
    assert "Resumo completo de exemplo." in document_text


def test_output_filename_number_increments_for_same_pdf(
    service: TaskService,
    pdf_path: Path,
) -> None:
    first_result = service.create_summary(pdf_path, "Qual é o assunto principal?")
    second_result = service.create_summary(pdf_path, "Qual é o assunto principal?")

    assert first_result.filename == "documento_1.docx"
    assert second_result.filename == "documento_2.docx"


def test_create_summary_fails_when_sentence_is_not_in_pdf(
    service: TaskService,
    repository: TaskRepository,
    pdf_path: Path,
) -> None:
    with pytest.raises(TaskProcessingError, match="A sentença não foi encontrada"):
        service.create_summary(pdf_path, "Uma sentença que não existe no PDF")

    tasks = repository.list()
    assert len(tasks) == 1
    assert tasks[0].status == "failed"


def test_guided_summary_uses_every_paragraph_containing_the_sentence(
    tmp_path: Path,
) -> None:
    class CapturingAdvisor:
        source_text = ""

        def summarize(self, source_text: str, sentence: str | None) -> str:
            self.source_text = source_text
            return "Resumo dos parágrafos pertinentes."

    pdf_path = tmp_path / "paragrafos.pdf"
    paragraphs = [
        "A inflação caiu após a política monetária adotada.",
        "O relatório também descreve o mercado de trabalho.",
        "A inflação caiu novamente no trimestre seguinte.",
    ]
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        for index, paragraph in enumerate(paragraphs):
            page.insert_text((72, 72 + index * 30), paragraph)
        pdf.save(pdf_path)

    advisor = CapturingAdvisor()
    service = TaskService(TaskRepository(), advisor)

    result = service.create_summary(pdf_path, "A inflação caiu")

    assert "A inflação caiu após a política monetária adotada." in advisor.source_text
    assert "A inflação caiu novamente no trimestre seguinte." in advisor.source_text
    assert "mercado de trabalho" not in advisor.source_text
    assert result.task.summary == "Resumo dos parágrafos pertinentes."


def test_get_task_returns_404_for_nonexistent_id(repository: TaskRepository) -> None:
    app = FastAPI()
    app.include_router(router)
    app.state.task_repository = repository

    with TestClient(app) as client:
        response = client.get("/tasks/id-inexistente")

    assert response.status_code == 404
    assert response.json() == {"detail": "Tarefa não encontrada."}
